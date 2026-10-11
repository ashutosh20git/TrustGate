"""PDF and Office attachments -> evidence items through the same splitter (v2).

PyMuPDF, python-docx, openpyxl and python-pptx extract text per page, sheet or
slide. Text a human reader would not see (invisible render mode, white or tiny
text, hidden runs, sheets, rows, columns and slides, off-page shapes) is kept
as a separate part and flagged "hidden_text". Metadata, comments and speaker
notes become their own parts. Any parser error turns the attachment into an
untrusted evidence item flagged "unparseable" (fail closed).

Images (attached, or embedded in PDF pages) are read too: QR codes are decoded
with OpenCV (quishing: a phishing link no text filter sees) and text is read
with Tesseract OCR. Both become untrusted evidence; QR links go to URL
reputation. A missing OCR or QR backend is flagged ("ocr_unavailable",
"qr_unavailable"), never silent.
"""

from __future__ import annotations

import hashlib
import io
import os
import re
import zipfile
from dataclasses import dataclass, field

MAX_ATTACHMENT_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
MAX_PAGES = 200
MAX_CELLS = 50_000
MAX_IMAGE_PIXELS = 40_000_000   # decompression-bomb guard
MAX_PDF_IMAGES = 20
_URL = re.compile(r"(?i)\bhttps?://[^\s<>\"']+")

EXTENSION_KIND = {
    ".pdf": "pdf",
    ".docx": "docx", ".docm": "docx",
    ".xlsx": "xlsx", ".xlsm": "xlsx",
    ".pptx": "pptx", ".pptm": "pptx",
    ".png": "image", ".jpg": "image", ".jpeg": "image", ".gif": "image", ".bmp": "image", ".webp": "image",
}
_WT = re.compile(r"<w:t(?:\s[^>]*)?>(.*?)</w:t>", re.S)
_ACTIVE_FORMULA = re.compile(r"\b(WEBSERVICE|HYPERLINK|IMPORTDATA|IMPORTXML|CALL|EXEC|REGISTER)\s*\(", re.I)


@dataclass
class AttachmentPart:
    field: str
    text: str
    flags: list[str] = field(default_factory=list)


@dataclass
class AttachmentResult:
    index: int
    filename: str            # attacker-controlled; becomes its own evidence item
    sha256: str
    kind: str
    parts: list[AttachmentPart] = field(default_factory=list)
    flags: list[str] = field(default_factory=list)
    parse_failed: bool = False
    hidden_content: bool = False
    has_macros: bool = False
    has_qr: bool = False
    urls: list[str] = field(default_factory=list)  # from QR codes: sent to URL reputation


class UnparseableAttachment(ValueError):
    pass


def detect_kind(data: bytes) -> str:
    """By magic bytes, never by the attacker's filename."""
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:4] == b"PK\x03\x04":
        try:
            names = zipfile.ZipFile(io.BytesIO(data)).namelist()
        except zipfile.BadZipFile:
            return "zip_corrupt"
        if "[Content_Types].xml" in names:
            for prefix, kind in (("word/", "docx"), ("xl/", "xlsx"), ("ppt/", "pptx")):
                if any(n.startswith(prefix) for n in names):
                    return kind
        return "zip"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "ole"  # legacy .doc/.xls/.ppt: not parsed, fails closed
    if (data[:8] == b"\x89PNG\r\n\x1a\n" or data[:3] == b"\xff\xd8\xff" or data[:4] == b"GIF8"
            or data[:2] == b"BM" or (data[:4] == b"RIFF" and data[8:12] == b"WEBP")):
        return "image"
    return "unknown"


def parse_attachment(index: int, filename: str, data: bytes) -> AttachmentResult:
    result = AttachmentResult(index=index, filename=filename or "", sha256=hashlib.sha256(data).hexdigest(),
                              kind=detect_kind(data))
    ext = os.path.splitext((filename or "").lower())[1]
    if ext in EXTENSION_KIND and EXTENSION_KIND[ext] != result.kind:
        result.flags.append("type_mismatch")
    try:
        if len(data) > MAX_ATTACHMENT_BYTES:
            raise UnparseableAttachment("oversize")
        if result.kind in ("docx", "xlsx", "pptx"):
            _check_zip(data, result)
        parser = _PARSERS.get(result.kind)
        if parser is None:
            raise UnparseableAttachment(f"unsupported type: {result.kind}")
        parser(data, result)
    except Exception as exc:  # any parser failure fails closed
        result.parse_failed = True
        result.flags.append("unparseable")
        result.parts = [AttachmentPart(
            field=f"attachment[{index}]",
            text=f"[unparseable attachment: {result.kind} ({type(exc).__name__})]",
            flags=["unparseable"],
        )]
    result.hidden_content = any("hidden_text" in p.flags for p in result.parts)
    return result


def _check_zip(data: bytes, result: AttachmentResult) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        infos = zf.infolist()
        total = sum(i.file_size for i in infos)
        compressed = max(1, sum(i.compress_size for i in infos))
        if total > MAX_UNCOMPRESSED_BYTES or total / compressed > MAX_COMPRESSION_RATIO:
            raise UnparseableAttachment("zip bomb guard")
        if any(i.filename.lower().endswith("vbaproject.bin") for i in infos):
            result.has_macros = True
            result.flags.append("macros")


def _add_part(result: AttachmentResult, suffix: str, chunks: list[str], flags: list[str]) -> None:
    text = "\n".join(c for c in chunks if c and c.strip())
    if text.strip():
        result.parts.append(AttachmentPart(f"attachment[{result.index}].{suffix}", text, list(flags)))


def _metadata_part(result: AttachmentResult, pairs: list[tuple[str, object]]) -> None:
    lines = [f"{k}: {v}" for k, v in pairs if v not in (None, "")]
    _add_part(result, "metadata", lines, ["metadata"])


# --- PDF (PyMuPDF) ----------------------------------------------------------

def _is_white(color) -> bool:
    if color is None:
        return False
    comps = list(color) if isinstance(color, (tuple, list)) else [color]
    if len(comps) == 4:  # CMYK: white is all zeros
        return all(c <= 0.05 for c in comps)
    return bool(comps) and all(c >= 0.95 for c in comps)


def _parse_pdf(data: bytes, result: AttachmentResult) -> None:
    import fitz  # PyMuPDF

    doc = fitz.open(stream=data, filetype="pdf")
    try:
        if doc.needs_pass:
            raise UnparseableAttachment("encrypted pdf")
        meta = doc.metadata or {}
        images_seen = 0
        _metadata_part(result, [(k, meta.get(k)) for k in ("title", "author", "subject", "keywords", "creator")])
        for xref in range(1, doc.xref_length()):
            obj = doc.xref_object(xref, compressed=False) or ""
            if "/JavaScript" in obj or "/JS " in obj or "/Launch" in obj:
                result.has_macros = True
                if "pdf_active_content" not in result.flags:
                    result.flags.append("pdf_active_content")
                break
        if doc.embfile_count() > 0:
            result.flags.append("embedded_file")
        for pno, page in enumerate(doc):
            if pno >= MAX_PAGES:
                result.flags.append("truncated")
                break
            visible: list[str] = []
            hidden: list[str] = []
            page_rect = page.rect
            for span in page.get_texttrace():
                text = "".join(chr(c[0]) for c in span.get("chars", ()) if c[0] > 0)
                if not text.strip():
                    continue
                bbox = fitz.Rect(span.get("bbox", (0, 0, 0, 0)))
                invisible = (
                    span.get("type") == 3
                    or float(span.get("opacity", 1.0)) <= 0.01
                    or float(span.get("size", 12)) < 1.0
                    or _is_white(span.get("color"))
                    or not page_rect.intersects(bbox)
                )
                (hidden if invisible else visible).append(text)
            _add_part(result, f"page[{pno}]", [" ".join(visible)], [])
            _add_part(result, f"page[{pno}].hidden", [" ".join(hidden)], ["hidden_text"])
            notes = [str((annot.info or {}).get("content", "")) for annot in (page.annots() or [])]
            _add_part(result, f"page[{pno}].annotations", notes, ["annotation"])
            for n, img in enumerate(page.get_images(full=True)):
                if images_seen >= MAX_PDF_IMAGES:
                    if "images_truncated" not in result.flags:
                        result.flags.append("images_truncated")
                    break
                images_seen += 1
                try:
                    raw = doc.extract_image(img[0]).get("image", b"")
                except Exception:  # noqa: BLE001 - one bad image does not void the page
                    continue
                if raw:
                    try:
                        _scan_image(raw, result, f"page[{pno}].image[{n}]")
                    except UnparseableAttachment:
                        _flag(result, "image_unreadable")  # e.g. JPEG 2000 / JBIG2 scans
    finally:
        doc.close()


# --- Images: QR codes (OpenCV) and OCR (Tesseract) ---------------------------------

def _flag(result: AttachmentResult, flag: str) -> None:
    if flag not in result.flags:
        result.flags.append(flag)


def _scan_image(data: bytes, result: AttachmentResult, label: str) -> None:
    try:
        import cv2
        import numpy as np
    except ImportError:
        _flag(result, "qr_unavailable")
        _flag(result, "ocr_unavailable")
        return
    img = cv2.imdecode(np.frombuffer(data, dtype=np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        raise UnparseableAttachment("image does not decode")
    if img.shape[0] * img.shape[1] > MAX_IMAGE_PIXELS:
        raise UnparseableAttachment("image too large")
    try:
        ok, texts, _points, _ = cv2.QRCodeDetector().detectAndDecodeMulti(img)
    except cv2.error:
        ok, texts = False, ()
    payloads = [t for t in (texts if ok else ()) if t]
    if payloads:
        result.has_qr = True
        _flag(result, "qr_code")
        for n, text in enumerate(payloads):
            _add_part(result, f"{label}.qr[{n}]", [text], ["qr_code"])
            result.urls.extend(_URL.findall(text))
    try:
        import pytesseract

        pytesseract.get_tesseract_version()
    except Exception:  # noqa: BLE001 - not installed: flagged, not silent
        _flag(result, "ocr_unavailable")
        return
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    _add_part(result, f"{label}.ocr", [pytesseract.image_to_string(gray, timeout=20)], ["ocr"])


def _parse_image(data: bytes, result: AttachmentResult) -> None:
    _scan_image(data, result, "image")


# --- DOCX (python-docx) -----------------------------------------------------

def _docx_run_hidden(run) -> bool:
    from docx.shared import Pt

    font = run.font
    if font.hidden:
        return True
    if font.size is not None and font.size <= Pt(1):
        return True
    try:
        if font.color is not None and font.color.type is not None and font.color.rgb is not None:
            return str(font.color.rgb).upper() == "FFFFFF"
    except (AttributeError, ValueError):
        return False
    return False


def _parse_docx(data: bytes, result: AttachmentResult) -> None:
    from docx import Document

    doc = Document(io.BytesIO(data))
    visible: list[str] = []
    hidden: list[str] = []

    def walk(paragraphs) -> None:
        for paragraph in paragraphs:
            seen, unseen = [], []
            for run in paragraph.runs:
                (unseen if _docx_run_hidden(run) else seen).append(run.text)
            visible.append("".join(seen))
            hidden.append("".join(unseen))

    walk(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                walk(cell.paragraphs)
    for section in doc.sections:
        walk(section.header.paragraphs)
        walk(section.footer.paragraphs)

    _add_part(result, "body", visible, [])
    _add_part(result, "body.hidden", hidden, ["hidden_text"])
    props = doc.core_properties
    _metadata_part(result, [("title", props.title), ("subject", props.subject), ("author", props.author),
                            ("comments", props.comments), ("keywords", props.keywords),
                            ("category", props.category), ("last_modified_by", props.last_modified_by)])
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        if "word/comments.xml" in zf.namelist():
            xml = zf.read("word/comments.xml").decode("utf-8", errors="replace")
            _add_part(result, "comments", [" ".join(_WT.findall(xml))], ["comment"])


# --- XLSX (openpyxl) --------------------------------------------------------

def _xlsx_white(cell) -> bool:
    color = getattr(cell.font, "color", None)
    rgb = getattr(color, "rgb", None)
    return isinstance(rgb, str) and rgb.upper().endswith("FFFFFF")


def _parse_xlsx(data: bytes, result: AttachmentResult) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    cells_seen = 0
    for ws in wb.worksheets:
        sheet_hidden = ws.sheet_state != "visible"
        hidden_rows = {r for r, dim in ws.row_dimensions.items() if dim.hidden}
        hidden_cols = {c for c, dim in ws.column_dimensions.items() if dim.hidden}
        visible: list[str] = []
        hidden: list[str] = []
        comments: list[str] = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.comment is not None and cell.comment.text:
                    comments.append(cell.comment.text)
                if cell.value is None:
                    continue
                cells_seen += 1
                if cells_seen > MAX_CELLS:
                    result.flags.append("truncated")
                    break
                text = str(cell.value)
                if text.startswith("=") and _ACTIVE_FORMULA.search(text):
                    if "active_formula" not in result.flags:
                        result.flags.append("active_formula")
                is_hidden = (sheet_hidden or cell.row in hidden_rows or cell.column_letter in hidden_cols
                             or _xlsx_white(cell))
                (hidden if is_hidden else visible).append(text)
        name = re.sub(r"[^A-Za-z0-9_-]", "_", ws.title)[:40]
        _add_part(result, f"sheet[{name}]", [" | ".join(visible)], [])
        _add_part(result, f"sheet[{name}].hidden", [" | ".join(hidden)], ["hidden_text"])
        _add_part(result, f"sheet[{name}].comments", comments, ["comment"])
    props = wb.properties
    _metadata_part(result, [("title", props.title), ("subject", props.subject), ("creator", props.creator),
                            ("description", props.description), ("keywords", props.keywords)])


# --- PPTX (python-pptx) -----------------------------------------------------

def _pptx_run_hidden(run) -> bool:
    from pptx.util import Pt

    font = run.font
    if font.size is not None and font.size <= Pt(1):
        return True
    try:
        if font.color is not None and font.color.type is not None and font.color.rgb is not None:
            return str(font.color.rgb).upper() == "FFFFFF"
    except (AttributeError, ValueError):
        return False
    return False


def _parse_pptx(data: bytes, result: AttachmentResult) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(io.BytesIO(data))
    width, height = prs.slide_width or 0, prs.slide_height or 0

    def off_slide(shape) -> bool:
        if shape.left is None or shape.top is None or not width or not height:
            return False
        right = shape.left + (shape.width or 0)
        bottom = shape.top + (shape.height or 0)
        return shape.left >= width or shape.top >= height or right <= 0 or bottom <= 0

    def walk(shapes, visible: list[str], hidden: list[str], force_hidden: bool) -> None:
        for shape in shapes:
            shape_hidden = force_hidden or off_slide(shape)
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes, visible, hidden, shape_hidden)
                continue
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        (hidden if shape_hidden or _pptx_run_hidden(run) else visible).append(run.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        (hidden if shape_hidden else visible).append(cell.text)

    for i, slide in enumerate(prs.slides):
        if i >= MAX_PAGES:
            result.flags.append("truncated")
            break
        slide_hidden = slide._element.get("show") == "0"
        visible: list[str] = []
        hidden: list[str] = []
        walk(slide.shapes, visible, hidden, slide_hidden)
        _add_part(result, f"slide[{i}]", [" ".join(visible)], [])
        _add_part(result, f"slide[{i}].hidden", [" ".join(hidden)], ["hidden_text"])
        if slide.has_notes_slide:
            _add_part(result, f"slide[{i}].notes", [slide.notes_slide.notes_text_frame.text], ["speaker_notes"])
    props = prs.core_properties
    _metadata_part(result, [("title", props.title), ("subject", props.subject), ("author", props.author),
                            ("comments", props.comments), ("keywords", props.keywords)])


_PARSERS = {"pdf": _parse_pdf, "docx": _parse_docx, "xlsx": _parse_xlsx, "pptx": _parse_pptx, "image": _parse_image}
